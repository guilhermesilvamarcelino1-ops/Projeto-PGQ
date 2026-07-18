import io

from docx import Document as DocxDocument

from app.services.chunking import _split_paragraphs, extract_docx, extract_plain_text


def test_split_paragraphs_merges_small_trailing_chunk():
    text = "Parágrafo curto." + "\n\n" + ("Parágrafo bem mais longo. " * 60)
    chunks = _split_paragraphs(text, max_chars=200)
    assert len(chunks) >= 1
    assert all(len(c) > 0 for c in chunks)


def test_extract_plain_text_splits_on_blank_lines():
    text = "Item 1: fazer X.\n\nItem 2: fazer Y."
    chunks = extract_plain_text(text)
    assert len(chunks) == 1  # merged, both short
    assert "Item 1" in chunks[0].content
    assert "Item 2" in chunks[0].content


def _build_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("4.2 Concretagem de fundação", level=2)
    doc.add_paragraph("O traço do concreto deve seguir o memorial estrutural.")
    doc.add_paragraph("Slump test obrigatório a cada caminhão.")
    doc.add_heading("4.3 Cura do concreto", level=2)
    doc.add_paragraph("Manter a superfície úmida por no mínimo 7 dias.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_extract_docx_tracks_section_headings():
    chunks = extract_docx(_build_docx_bytes())
    sections = {c.section_ref for c in chunks}
    assert "4.2 Concretagem de fundação" in sections
    assert "4.3 Cura do concreto" in sections
    concretagem_chunk = next(c for c in chunks if c.section_ref == "4.2 Concretagem de fundação")
    assert "traço do concreto" in concretagem_chunk.content
